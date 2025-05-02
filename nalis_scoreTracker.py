
import tkinter as tk
from tkinter import *
from tkinter import messagebox
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook




window = tk.Tk()
window.geometry("350x250")
window.title("Score Tracker")
label_font = ("Arial", 12)
entry_bg = "#ffffff"
entry_fg = "#333333"
button_bg = "#4CAF50"
button_fg = "#ffffff"
button2_bg = "#f44336"

l1 = Label(window, text="Name", font=label_font, bg="#f0f4f7")
l2 = Label(window, text="Grade", font=label_font, bg="#f0f4f7")

l1.grid(row=0, column=0, padx=10, pady=10, sticky= E)
l2.grid(row=1, column=0, padx=10, pady=10, sticky= E)

name_entry = tk.Entry(window, width=40)
grade_entry = tk.Entry(window, width=40)


name_entry.grid(row=0, column=1, pady=5)
grade_entry.grid(row=1, column=1, pady=5)








def inputs():
    name = name_entry.get()
    grade_str = grade_entry.get()

    if not name or not grade_str:
        messagebox.showerror("Please Fill in all Textfields")
        return False

    try:
        grade = int(grade_entry.get())
    except ValueError:
        messagebox.showerror("Grade must be a valid number")
        return False

    return True


    


def excel():
    wb = load_workbook("Student_scores.xlsx")
    ws = wb["Student_scores"]


    for cell in ws[1]:
        cell.font = Font(bold=True)


    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max_length + 2

    wb.save("Student_scores.xlsx")





def grade_validation(grade):
    if grade > 100:
        return "Invalid grade"
    elif grade >= 75:
        return "Passed"
    else:
        return "Failed"


def excel_save():

    if not inputs():
        return

    name = name_entry.get()
    grade = int(grade_entry.get())

    wb = load_workbook("Student_scores.xlsx")
    ws = wb["Student_scores"]
    ws.append([name, grade, grade_validation(grade)])
    wb.save("Student_scores.xlsx")

    excel()
    messagebox.showinfo("Success", "Data saved successfully!")

    name_entry.delete(0, tk.END)
    grade_entry.delete(0, tk.END)

    
def clear():
    name_entry.delete(0, tk.END)
    grade_entry.delete(0, tk.END)



def show_data():
    wb = load_workbook("Student_scores.xlsx")
    ws = wb["Student_scores"]

    data_window = tk.Toplevel(window)
    data_window.title("Student_scores")
   

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        for j, value in enumerate(row):
            label = tk.Label(data_window, text=value, borderwidth=1, relief="solid", padx=6, pady=3)
            label.grid(row=i, column=j)

b1 = Button(window, text="Save",bg=button_bg, fg=button_fg, font=label_font, command=excel_save, width=10)
b2 = Button(window, text="View Database",font=label_font, command=show_data, width=15)
b3 = Button(window, text="clear", bg=button2_bg, fg=button_fg, font=label_font,command=clear, width=10)


b1.grid(row=3, column=0, columnspan=2, pady=10)
b2.grid(row=4, column=0, columnspan=2, pady=5)
b3.grid(row=5, column=0, columnspan=2, pady=5)

window.mainloop()